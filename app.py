from fastapi import FastAPI, HTTPException, Request
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from pathlib import Path
from contextlib import contextmanager
from datetime import datetime, timezone
import sqlite3,json,io,os,math
import openpyxl
ROOT=Path(__file__).parent
DB=Path(os.environ.get('PLANNER_DB',str(ROOT/'data/planner.sqlite3')))
STATIONS={'auto':'Automat','lak':'Półautomat · lakiery','utw':'Półautomat · utwardzacze','stock':'Stan / etykietowanie'}
app=FastAPI(title='EXLAK • Planowanie produkcji')
@contextmanager
def db():
 c=sqlite3.connect(DB,timeout=15);c.row_factory=sqlite3.Row;c.execute('PRAGMA foreign_keys=ON')
 try:
  c.execute('BEGIN IMMEDIATE');yield c;c.commit()
 except: c.rollback();raise
 finally:c.close()
def now():return datetime.now(timezone.utc).isoformat()
def init():
 DB.parent.mkdir(parents=True,exist_ok=True)
 with db() as c:
  c.executescript('''CREATE TABLE IF NOT EXISTS orders(id INTEGER PRIMARY KEY,client TEXT NOT NULL,order_no TEXT NOT NULL,product TEXT NOT NULL,recipe TEXT NOT NULL,qty INTEGER NOT NULL CHECK(qty>0),volume REAL NOT NULL CHECK(volume>0),code TEXT,comment TEXT,due TEXT,priority TEXT DEFAULT 'normal',source TEXT,created TEXT);CREATE TABLE IF NOT EXISTS jobs(id INTEGER PRIMARY KEY,order_id INTEGER NOT NULL REFERENCES orders(id),station TEXT NOT NULL,qty INTEGER NOT NULL CHECK(qty>0),done INTEGER DEFAULT 0 CHECK(done>=0 AND done<=qty),state TEXT DEFAULT 'draft',note TEXT DEFAULT '',position INTEGER,created TEXT);CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY,job_id INTEGER REFERENCES jobs(id),action TEXT,detail TEXT,created TEXT);''')
init()
def event(c,j,a,d=''):c.execute('INSERT INTO events(job_id,action,detail,created) VALUES(?,?,?,?)',(j,a,d,now()))
def insert_orders(c,rows,source):
 count=0
 for r in rows:
  qty=float(r['qty']);vol=float(str(r['volume']).replace(',','.'))
  if not math.isfinite(qty) or not math.isfinite(vol) or qty<=0 or vol<=0 or qty!=int(qty):raise ValueError('Ilość musi być dodatnią liczbą całkowitą, a objętość dodatnią liczbą.')
  if not r.get('order') or not r.get('client') or not r.get('product'):raise ValueError('Brak klienta, numeru zamówienia lub produktu.')
  # A repeated source is rejected as a whole; identical rows within one file are retained.
  c.execute('INSERT INTO orders(client,order_no,product,recipe,qty,volume,code,comment,source,created) VALUES(?,?,?,?,?,?,?,?,?,?)',(r['client'],r['order'],r['product'],r.get('recipe') or '',int(qty),vol,r.get('code') or '',r.get('comment') or '',source,now()));count+=1
 return count
@app.get('/api/state')
def state():
 with db() as c:
  orders=[dict(r) for r in c.execute('SELECT o.*,COALESCE((SELECT SUM(qty) FROM jobs WHERE order_id=o.id),0) allocated FROM orders o ORDER BY id')]
  jobs=[dict(r) for r in c.execute('SELECT j.*,o.client,o.order_no,o.product,o.recipe,o.volume,o.code,o.comment,o.due,o.priority FROM jobs j JOIN orders o ON o.id=j.order_id ORDER BY j.position,j.id')]
  events=[dict(r) for r in c.execute('SELECT * FROM events ORDER BY id DESC LIMIT 30')]
 return dict(orders=orders,jobs=jobs,events=events,stations=STATIONS)
@app.post('/api/demo')
def demo():
 with db() as c:
  if c.execute('SELECT COUNT(*) FROM orders').fetchone()[0]:raise HTTPException(409,'Dane są już w bazie. Import demonstracyjny działa tylko w pustej bazie.')
  count=insert_orders(c,json.loads((ROOT/'data/demo-orders.json').read_text(encoding='utf-8')),'Przykład · 24.08.2026')
 return {'count':count}
@app.post('/api/import')
async def upload(request:Request):
 raw=await request.body()
 if len(raw)>8*1024*1024:raise HTTPException(413,'Maksymalny rozmiar pliku to 8 MB.')
 import hashlib
 source='xlsx:'+hashlib.sha256(raw).hexdigest()
 try:
  w=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
  s=w.active
  if s.max_row>10000 or s.max_column>100:raise ValueError('Arkusz jest zbyt duży.')
  it=s.iter_rows(values_only=True);headers=next(it)
  if [str(x or '').strip().upper() for x in headers[:5]]!=['NAZWA KONTRAHENTA','NUMER ZAM.','NAZWA TOWARU','NAZWA RECEPTURY','ILOŚĆ']:raise ValueError('Wybierz plik „Zestawienie do planu produkcyjnego”, nie plan maszyn.')
  rows=[]
  for i,r in enumerate(it,2):
   if not any(x is not None for x in r):continue
   if len(r)<9:raise ValueError(f'Niepełny wiersz {i}.')
   rows.append(dict(client=r[0],order=r[1],product=r[2],recipe=r[3],qty=r[4],code=r[6],volume=r[7],comment=r[9] if len(r)>9 else ''))
  if not rows:raise ValueError('Plik nie zawiera pozycji zamówień.')
  with db() as c:
   if c.execute('SELECT 1 FROM orders WHERE source=?',(source,)).fetchone():raise HTTPException(409,'Ten plik był już importowany.')
   existing={r[0] for r in c.execute('SELECT DISTINCT order_no FROM orders')}
   overlap=existing & {str(r['order']) for r in rows}
   if overlap:raise HTTPException(409,'Plik zawiera zamówienia obecne w bazie. Kolejne zestawienia mogą być aktualizacjami — import przerwany, aby uniknąć podwójnego planowania.')
   count=insert_orders(c,rows,source)
  return {'count':count}
 except HTTPException:raise
 except Exception as e:raise HTTPException(400,str(e) if isinstance(e,ValueError) else 'Nie można odczytać pliku XLSX.')
class Allocation(BaseModel):
 order_id:int;station:str;qty:int=Field(gt=0);note:str=Field(default='',max_length=2000)
@app.post('/api/jobs')
def allocate(x:Allocation):
 if x.station not in STATIONS:raise HTTPException(400,'Nieznane stanowisko.')
 with db() as c:
  o=c.execute('SELECT *,COALESCE((SELECT SUM(qty) FROM jobs WHERE order_id=orders.id),0) allocated FROM orders WHERE id=?',(x.order_id,)).fetchone()
  if not o:raise HTTPException(404,'Nie znaleziono pozycji.')
  if x.qty>o['qty']-o['allocated']:raise HTTPException(409,'Ilość przekracza pozostałą ilość zamówienia.')
  pos=c.execute('SELECT COALESCE(MAX(position),0)+1 FROM jobs').fetchone()[0]
  j=c.execute('INSERT INTO jobs(order_id,station,qty,note,position,created) VALUES(?,?,?,?,?,?)',(x.order_id,x.station,x.qty,x.note,pos,now())).lastrowid;event(c,j,'Przydzielono',STATIONS[x.station])
 return {'id':j}
@app.delete('/api/jobs/{jid}')
def remove(jid:int):
 with db() as c:
  j=c.execute('SELECT * FROM jobs WHERE id=?',(jid,)).fetchone()
  if not j or j['state']!='draft':raise HTTPException(409,'Usuwać można tylko zadania robocze.')
  c.execute('DELETE FROM events WHERE job_id=?',(jid,));c.execute('DELETE FROM jobs WHERE id=?',(jid,))
 return {'ok':True}
@app.post('/api/publish')
def publish():
 with db() as c:
  ids=[r[0] for r in c.execute("SELECT id FROM jobs WHERE state='draft'")]
  c.execute("UPDATE jobs SET state='ready' WHERE state='draft'")
  for j in ids:event(c,j,'Opublikowano')
 return {'count':len(ids)}
class Meta(BaseModel):
 due:str='';priority:str='normal'
@app.patch('/api/orders/{oid}')
def meta(oid:int,x:Meta):
 if x.priority not in ['normal','high']:raise HTTPException(400,'Nieprawidłowy priorytet.')
 if x.due:
  try:datetime.strptime(x.due,'%Y-%m-%d')
  except ValueError:raise HTTPException(400,'Nieprawidłowy termin.')
 with db() as c:c.execute('UPDATE orders SET due=?,priority=? WHERE id=?',(x.due,x.priority,oid))
 return {'ok':True}
class Action(BaseModel):
 action:str;qty:int=Field(default=0,ge=0);note:str=Field(default='',max_length=2000)
@app.post('/api/jobs/{jid}/action')
def action(jid:int,x:Action):
 with db() as c:
  j=c.execute('SELECT * FROM jobs WHERE id=?',(jid,)).fetchone()
  if not j:raise HTTPException(404,'Nie znaleziono zadania.')
  status=j['state'];done=j['done'];target=status
  if x.action=='start' and status in ['ready','paused']:
   if c.execute("SELECT 1 FROM jobs WHERE station=? AND state='running'",(j['station'],)).fetchone():raise HTTPException(409,'Na stanowisku trwa inne zadanie. Najpierw je wstrzymaj lub zakończ.')
   target='running'
  elif x.action=='pause' and status=='running':
   if not x.note.strip():raise HTTPException(400,'Podaj przyczynę wstrzymania.')
   target='paused'
  elif x.action=='report' and status=='running':
   if not x.qty or done+x.qty>j['qty']:raise HTTPException(400,'Podaj dodatnią ilość nieprzekraczającą ilości pozostałej.')
   done+=x.qty
   if done==j['qty']:target='done'
  else:raise HTTPException(409,'Ta operacja nie jest dostępna dla aktualnego statusu zadania.')
  c.execute('UPDATE jobs SET state=?,done=? WHERE id=?',(target,done,jid));event(c,jid,x.action,json.dumps({'qty':x.qty,'note':x.note},ensure_ascii=False))
 return {'ok':True}
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')
@app.get('/')
def index():return FileResponse(ROOT/'static/index.html')
